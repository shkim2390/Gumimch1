# -*- coding: utf-8 -*-
import os
import sys
import math
from datetime import datetime, date
from pathlib import Path
import streamlit as st
import pandas as pd
import openpyxl

# ==================================================
# 경로 및 파일 설정
# ==================================================
BASE_DIR = Path(__file__).resolve().parent
ID_FILE = BASE_DIR / "ID.xlsx"
CONFIG_FILE = BASE_DIR / "Config.txt"
LOG_FILE = BASE_DIR / "Log.txt"
UPLOAD_DIR = BASE_DIR / "업로드"
UPLOAD_SAVE_DIR = UPLOAD_DIR / "save"
ACCOUNT_FILE = BASE_DIR / "IDPW.txt"

# 기본 설정값
DEFAULT_NAMES = ["강주경", "권해선", "김소원", "김점수", "임영오", "전현숙", "이태권", "조인제"]
DEFAULT_TYPES = [
    "십일조", "주일헌금", "영아부헌금", "유치부헌금", "아동부헌금",
    "청소년부헌금", "청년부헌금", "감사헌금", "부활절헌금", "맥추절헌금",
    "추수절헌금", "성탄절헌금", "신년헌금", "월정헌금", "심방헌금",
    "선교헌금", "부흥회헌금", "일천번제헌금", "일반건축헌금",
    "임마누엘헌금", "이웃사랑헌금", "꽃꽂이헌금", "목적헌금", "차량헌금",
    "이월금", "금융수입", "기타수입"
]
DEFAULT_UNITS = ["x(원)", "x1000(원)", "x10000(원)"]
DEFAULT_WORSHIP_TIMES = ["", "09:00", "11:00", "12:00"]
DEFAULT_PERSON = ""
DEFAULT_TYPE = ""

def load_config():
    global DEFAULT_NAMES, DEFAULT_TYPES, DEFAULT_UNITS, DEFAULT_WORSHIP_TIMES, DEFAULT_PERSON, DEFAULT_TYPE
    if not CONFIG_FILE.exists():
        return
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                key, val = line.split("=", 1)
                key = key.strip()
                val_str = val.strip()
                items = [item.strip() for item in val_str.split(",") if item.strip()]

                if key == "NAMES" and items:
                    DEFAULT_NAMES = items
                elif key == "TYPES" and items:
                    DEFAULT_TYPES = items
                elif key == "UNITS" and items:
                    DEFAULT_UNITS = items
                elif key == "WORSHIP_TIMES" and items:
                    DEFAULT_WORSHIP_TIMES = items
                elif key == "DEFAULT_PERSON":
                    DEFAULT_PERSON = val_str
                elif key == "DEFAULT_TYPE":
                    DEFAULT_TYPE = val_str
    except Exception:
        pass

load_config()

# ==================================================
# 유틸리티 함수
# ==================================================
def normalize_lookup_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, (int, float)):
        numeric = float(value)
        if numeric.is_integer():
            return str(int(numeric))
        return format(numeric, "g")
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            numeric = float(text)
            if numeric.is_integer():
                return str(int(numeric))
        except ValueError:
            pass
    return text

def donation_category(kind):
    general_types = {"십일조", "주일헌금", "영아부헌금", "유치부헌금", "아동부헌금", "청소년부헌금", "청년부헌금"}
    thanks_types = {"감사헌금"}
    season_types = {"부활절헌금", "맥추절헌금", "추수절헌금"}
    special_types = {"신년헌금", "월정헌금", "심방헌금", "선교헌금", "부흥회헌금", "일천번제헌금", "임마누엘헌금", "이웃사랑헌금", "꽃꽂이헌금", "목적헌금", "차량헌금", "성탄절헌금", "일반건축헌금"}
    other_types = {"이월금", "금융수입", "기타수입"}

    if kind in general_types: return "일반계정" # 출력 양식에 맞춤
    if kind in thanks_types: return "감사헌금"
    if kind in season_types: return "절기헌금"
    if kind in special_types: return "특별헌금"
    if kind in other_types: return "기타수입"
    return "일반계정"

def month_week_number(date_value):
    if isinstance(date_value, datetime):
        day = date_value.day
    else:
        try:
            day = datetime.strptime(str(date_value).strip(), "%Y-%m-%d").day
        except:
            day = 1
    return int(math.ceil(day / 7.0))

# ==================================================
# ID.xlsx 캐싱 로드
# ==================================================
@st.cache_data(ttl=60)
def load_id_excel():
    id_to_name = {}
    name_to_id = {}
    records = []
    headers = []

    if not ID_FILE.exists():
        return id_to_name, name_to_id, records, headers

    try:
        wb = openpyxl.load_workbook(ID_FILE, data_only=True)
        ws = wb.active if "ID" not in wb.sheetnames else wb["ID"]
        
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            headers = [str(cell) if cell is not None else "" for cell in rows[0][:7]]
            for row in rows[1:]:
                if not any(row):
                    continue
                name_val = normalize_lookup_value(row[0])
                id_val = normalize_lookup_value(row[1]) if len(row) > 1 else ""
                if name_val and id_val:
                    name_to_id[name_val] = id_val
                    id_to_name[id_val] = name_val
                records.append([normalize_lookup_value(cell) if idx in (0,1,3) else (str(cell) if cell is not None else "") for idx, cell in enumerate(row[:7])])
    except Exception as e:
        st.error(f"ID.xlsx 파일을 읽는 중 오류 발생: {e}")

    return id_to_name, name_to_id, records, headers

id_to_name, name_to_id, id_records, id_headers = load_id_excel()

# ==================================================
# Log.txt 파일 관리 (세션 스테이트 연동)
# ==================================================
def load_log_to_session():
    data = []
    if LOG_FILE.exists():
        try:
            with open(LOG_FILE, "r", encoding="utf-8") as f:
                for line in f:
                    parts = line.strip().split("\t")
                    if len(parts) >= 6:
                        data.append({
                            "kind": parts[0],
                            "id": parts[1],
                            "name": parts[2],
                            "amount": parts[3],
                            "spouse": parts[4],
                            "memo": parts[5]
                        })
        except Exception:
            pass
    return data

def save_session_to_log():
    try:
        with open(LOG_FILE, "w", encoding="utf-8") as f:
            for item in st.session_state.records:
                f.write(f"{item['kind']}\t{item['id']}\t{item['name']}\t{item['amount']}\t{item['spouse']}\t{item['memo']}\n")
    except Exception as e:
        st.error(f"Log.txt 저장 오류: {e}")

if "records" not in st.session_state:
    st.session_state.records = load_log_to_session()

# ==================================================
# Streamlit UI 구성 (아이패드 최적화 반응형 레이아웃)
# ==================================================
st.set_page_config(page_title="구미제일감리교회 재무부", layout="wide")
st.title("⛪ 구미제일감리교회 헌금 입력 시스템")

# 상단 입력 폼 영역
with st.form("donation_form", clear_on_submit=False):
    st.subheader("📝 헌금 정보 입력")
    
    col1, col2, col3, col4, col5 = st.columns([2, 2, 2, 1.5, 1.5])
    
    with col1:
        default_person_idx = DEFAULT_NAMES.index(DEFAULT_PERSON) if DEFAULT_PERSON in DEFAULT_NAMES else 0
        person = st.selectbox("입력자", DEFAULT_NAMES, index=default_person_idx)
    with col2:
        default_type_idx = DEFAULT_TYPES.index(DEFAULT_TYPE) if DEFAULT_TYPE in DEFAULT_TYPES else 0
        kind = st.selectbox("헌금종류", DEFAULT_TYPES, index=default_type_idx)
    with col3:
        date_val = st.date_input("날짜", value=datetime.now())
    with col4:
        unit = st.selectbox("단위", DEFAULT_UNITS, index=1) # 기본 x1000(원)
    with col5:
        worship_time = st.selectbox("예배시간", DEFAULT_WORSHIP_TIMES, index=DEFAULT_WORSHIP_TIMES.index("11:00") if "11:00" in DEFAULT_WORSHIP_TIMES else 0)

    col_id, col_name, col_amt, col_spouse, col_memo = st.columns([1.5, 2, 2, 2, 3.5])
    
    with col_id:
        input_id = st.text_input("ID", placeholder="ID 입력")
    with col_name:
        input_name = st.text_input("이름", placeholder="이름 입력")
    with col_amt:
        input_amount = st.text_input("금액", placeholder="숫자만 입력")
    with col_spouse:
        input_spouse = st.text_input("부부", placeholder="배우자 이름")
    with col_memo:
        input_memo = st.text_input("메모", placeholder="비고 및 메모")

    submitted = st.form_submit_button("➕ 입력 추가", use_container_width=True)
    
    if submitted:
        if not input_amount.strip():
            st.warning("금액을 입력해주세요.")
        else:
            try:
                clean_amt = float(input_amount.replace(",", ""))
                multiplier = {"x(원)": 1, "x1000(원)": 1000, "x10000(원)": 10000}.get(unit, 1)
                final_amt = clean_amt * multiplier
                
                # ID로 이름 자동완성 또는 역방향 체크
                resolved_id = normalize_lookup_value(input_id)
                resolved_name = input_name.strip()
                
                if resolved_id and not resolved_name and resolved_id in id_to_name:
                    resolved_name = id_to_name[resolved_id]
                elif resolved_name and not resolved_id and resolved_name in name_to_id:
                    resolved_id = name_to_id[resolved_name]

                st.session_state.records.append({
                    "kind": kind,
                    "id": resolved_id,
                    "name": resolved_name,
                    "amount": f"{int(final_amt):,}",
                    "spouse": input_spouse.strip(),
                    "memo": input_memo.strip()
                })
                save_session_to_log()
                st.success("성공적으로 추가되었습니다!")
                st.rerun()
            except ValueError:
                st.error("금액은 숫자로 정확히 입력해주세요.")

# ==================================================
# 이름 검색 팝업 / 섹션
# ==================================================
with st.expander("🔍 교인 이름 / ID 조회 (검색)"):
    search_query = st.text_input("검색할 이름 또는 ID를 입력하세요", key="search_box")
    if search_query:
        matched_results = [row for row in id_records if any(search_query in str(cell) for cell in row)]
        if matched_results:
            df_search = pd.DataFrame(matched_results, columns=id_headers[:len(matched_results[0])] if id_headers else None)
            st.dataframe(df_search, use_container_width=True)
        else:
            st.info("검색 결과가 없습니다.")

st.divider()

# ==================================================
# 입력 내역 확인 및 삭제 테이블
# ==================================================
st.subheader("📊 현재 입력 내역")

if st.session_state.records:
    df_display = pd.DataFrame(st.session_state.records)
    df_display.index = range(1, len(df_display) + 1)
    df_display.columns = ["헌금종류", "ID", "이름", "금액", "부부", "메모"]
    
    st.dataframe(df_display, use_container_width=True)

    # 합계 계산
    total_sum = sum([float(str(r["amount"]).replace(",", "")) for r in st.session_state.records if str(r["amount"]).replace(",", "").replace(".", "").isdigit()])
    st.markdown(f"**📌 총 인원:** {len(st.session_state.records)}명 &nbsp;&nbsp;&nbsp;&nbsp;|&nbsp;&nbsp;&nbsp;&nbsp; **💰 총 합계:** `{int(total_sum):,`원**")

    # 행 삭제 기능
    col_del1, col_del2 = st.columns([2, 5])
    with col_del1:
        del_idx = st.number_input("삭제할 행 번호", min_value=1, max_value=len(st.session_state.records), step=1)
        if st.button("🗑️ 선택 행 삭제", use_container_width=True):
            removed = st.session_state.records.pop(del_idx - 1)
            save_session_to_log()
            st.success(f"{del_idx}번째 행({removed['name']})이 삭제되었습니다.")
            st.rerun()
else:
    st.info("입력된 내역이 없습니다.")

st.divider()

# ==================================================
# 하단 액션 버튼 그룹 (저장, 확인, 업로드 등)
# ==================================================
col_b1, col_b2, col_b3, col_b4 = st.columns(4)

with col_b1:
    if st.button("🔄 데이터 불러오기 (CALL)", use_container_width=True):
        st.session_state.records = load_log_to_session()
        st.rerun()

with col_b2:
    if st.button("💾 ① 저장 및 엑셀 변환", use_container_width=True):
        if not st.session_state.records:
            st.warning("저장할 내역이 없습니다.")
        else:
            try:
                UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
                file_name = f"{kind}_{person}_{date_val.strftime('%Y%m%d')}.xlsx"
                for char in '<>:"/\\|?*': file_name = file_name.replace(char, "_")
                dest_path = UPLOAD_DIR / file_name

                # 엑셀 출력 생성 (openpyxl 활용)
                wb_out = openpyxl.Workbook()
                ws_out = wb_out.active
                ws_out.title = "출력"

                output_headers = [
                    "일자", "헌금자ID", "헌금자", "생년월일", "휴대전화", "은행",
                    "과목1", "과목2", "과목3", "과목4", "금액", "배우자", "메모",
                    "주", "예배 시간", "계전자", "계전 파트"
                ]
                ws_out.append(output_headers)

                week_no = month_week_number(date_val)
                excel_date_value = date_val.strftime("%Y-%m-%d")

                for item in st.session_state.records:
                    d_id = item["id"]
                    try: d_id_val = int(float(d_id))
                    except: d_id_val = d_id

                    try: amt_val = float(str(item["amount"]).replace(",", ""))
                    except: amt_val = 0

                    row_data = [
                        excel_date_value, d_id_val, item["name"], "", "", "",
                        "일반계정", donation_category(item["kind"]), item["kind"], "",
                        amt_val, item["spouse"], item["memo"], week_no, worship_time, person, ""
                    ]
                    ws_out.append(row_data)

                wb_out.save(dest_path)
                st.success(f"엑셀 파일이 성공적으로 생성되었습니다: {file_name}")
            except Exception as e:
                st.error(f"저장 중 오류 발생: {e}")

with col_b3:
    # 3업로드 폴더 기능 대신 최근 생성된 엑셀 파일 다운로드 제공
    if UPLOAD_DIR.exists():
        excel_files = sorted([f for f in UPLOAD_DIR.iterdir() if f.suffix == ".xlsx"], key=lambda x: x.stat().st_mtime, reverse=True)
        if excel_files:
            latest_file = excel_files[0]
            with open(latest_file, "rb") as f:
                st.download_button(
                    label="📥 ② 최신 엑셀 다운로드",
                    data=f,
                    file_name=latest_file.name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
        else:
            st.button("📥 ② 최신 엑셀 다운로드", disabled=True, use_container_width=True)

with col_b4:
    if st.button("🧹 전체 초기화", use_container_width=True):
        st.session_state.records = []
        if LOG_FILE.exists():
            LOG_FILE.unlink()
        st.success("초기화되었습니다.")
        st.rerun()
